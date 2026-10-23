"""Export the current Markdown GDD using the existing Word document styles."""

from pathlib import Path
import re
import sys
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, RGBColor
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT


def inline(paragraph, text):
    pattern = r"(\[[^\]]+\]\([^)]+\)|\*\*[^*]+\*\*|`[^`]+`)"
    for part in re.split(pattern, text):
        link = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", part)
        if link:
            label, target = link.groups()
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("r:id"), paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True))
            run = OxmlElement("w:r")
            props = OxmlElement("w:rPr")
            color = OxmlElement("w:color")
            color.set(qn("w:val"), "245B76")
            props.append(color)
            run.append(props)
            node = OxmlElement("w:t")
            node.text = label
            run.append(node)
            hyperlink.append(run)
            paragraph._p.append(hyperlink)
        elif part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        else:
            paragraph.add_run(part)


def table_style(table, width):
    table.autofit = False
    count = len(table.columns)
    ratios = {2: [0.27, 0.73], 3: [0.20, 0.41, 0.39], 4: [0.13, 0.24, 0.31, 0.32]}.get(count, [1 / count] * count)
    for column, ratio in zip(table.columns, ratios):
        column.width = int(width * ratio)
    borders = OxmlElement("w:tblBorders")
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        border = OxmlElement("w:" + edge)
        for key, value in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
            border.set(qn("w:" + key), value)
        borders.append(border)
    table._tbl.tblPr.append(borders)
    for row_index, row in enumerate(table.rows):
        row_props = row._tr.get_or_add_trPr()
        row_props.append(OxmlElement("w:cantSplit"))
        if row_index == 0:
            row_props.append(OxmlElement("w:tblHeader"))
        for cell, ratio in zip(row.cells, ratios):
            cell.width = int(width * ratio)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            margin = OxmlElement("w:tcMar")
            for edge, value in [("top", "85"), ("bottom", "85"), ("left", "95"), ("right", "95")]:
                node = OxmlElement("w:" + edge)
                node.set(qn("w:w"), value)
                node.set(qn("w:type"), "dxa")
                margin.append(node)
            props.append(margin)
            if row_index == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "E8EEF1")
                props.append(shade)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.space_before = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.12
                paragraph.paragraph_format.keep_with_next = row_index == 0
                for run in paragraph.runs:
                    run.font.size = Pt(10.5)
                    if row_index == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(0, 0, 0)


def export(source, template, output):
    document = Document(template)
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Heading 4", "List Bullet"]:
        style = document.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt({"Title": 22, "Subtitle": 12, "Heading 1": 16, "Heading 2": 13, "Heading 3": 11.5, "Heading 4": 11}.get(name, 11))
        style.paragraph_format.page_break_before = False
        style.paragraph_format.widow_control = True
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.line_spacing = 1.16
        if name.startswith("Heading"):
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.space_before = Pt(12)
    section = document.sections[0]
    width = section.page_width - section.left_margin - section.right_margin
    lines = source.read_text(encoding="utf-8-sig").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line or line == "---":
            index += 1
            continue
        if line.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                raw = lines[index].strip().strip("|")
                cells = [cell.strip() for cell in raw.split("|")]
                if not all(re.fullmatch(r":?-+:?", cell.replace(" ", "")) for cell in cells):
                    rows.append(cells)
                index += 1
            table = document.add_table(rows=0, cols=len(rows[0]))
            for row in rows:
                cells = table.add_row().cells
                for cell, text in zip(cells, row):
                    inline(cell.paragraphs[0], text)
            table_style(table, width)
            document.add_paragraph().paragraph_format.space_after = Pt(2)
            continue
        heading = re.match(r"^(#{1,5})\s+(.+)$", line)
        if heading:
            level, text = len(heading[1]), heading[2]
            paragraph = document.add_paragraph(style="Title" if level == 1 else f"Heading {min(level - 1, 4)}")
            inline(paragraph, text.replace("《", "").replace("》", "") if level == 1 else text)
        else:
            if line.startswith(">"):
                line = line.lstrip("> ")
            bullet = line.startswith("- ")
            if bullet:
                line = line[2:]
            paragraph = document.add_paragraph(style="List Bullet" if bullet else "Normal")
            inline(paragraph, line)
        index += 1
    document.core_properties.title = "未名之海游戏设计文档"
    document.core_properties.author = "Codex"
    document.core_properties.subject = "第一人称方向与小团队 MVP 设计 2026-10-20"
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(f"Saved {output}; paragraphs={len(document.paragraphs)}, tables={len(document.tables)}")


if __name__ == "__main__":
    export(*(Path(argument) for argument in sys.argv[1:4]))
